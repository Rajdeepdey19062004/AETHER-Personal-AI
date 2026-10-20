"""Document text extraction and chunking helpers for AETHER."""

import io
import re
from pathlib import Path

from docx import Document
from pypdf import PdfReader

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
DOCUMENT_DIR = DATA_DIR / "documents"
DOCUMENT_DIR.mkdir(parents=True, exist_ok=True)


def extract_text(filename: str, content: bytes) -> str:
    """Extract text from a PDF, DOCX, TXT, or Markdown document."""
    extension = Path(filename).suffix.lower()

    if extension == ".pdf":
        pdf = PdfReader(io.BytesIO(content))
        return "\n".join(page.extract_text() or "" for page in pdf.pages)

    if extension in {".txt", ".md"}:
        return content.decode("utf-8", errors="replace")

    if extension == ".docx":
        document = Document(io.BytesIO(content))
        sections = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        # A lot of reports and papers put their abstract, results, and tables in
        # Word tables; include those cells in the searchable text too.
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    sections.append(" | ".join(cells))
        return "\n".join(sections)

    raise ValueError("Supported formats: PDF, TXT, MD, DOCX")


def split_text(text: str, chunk_size: int = 900, overlap: int = 150) -> list[str]:
    """Split text into readable, overlapping chunks."""
    if chunk_size <= 0:
        raise ValueError("chunk_size must be greater than zero")
    if overlap < 0 or overlap >= chunk_size:
        raise ValueError("overlap must be non-negative and smaller than chunk_size")

    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        return []

    chunks = []
    start = 0

    while start < len(text):
        end = min(start + chunk_size, len(text))

        if end < len(text):
            boundary = text.rfind(" ", start, end)
            if boundary > start + chunk_size // 2:
                end = boundary

        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)

        if end >= len(text):
            break

        start = max(end - overlap, start + 1)

    return chunks
